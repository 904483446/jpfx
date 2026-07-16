from io import BytesIO

import pytest
from app.core.errors import AppError
from app.services.document_parser import extract_text
from docx import Document as DocxDocument


def make_pdf_with_text(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(f'BT /F1 12 Tf 20 100 Td ({text}) Tj ET')} >>\n"
        f"stream\nBT /F1 12 Tf 20 100 Td ({text}) Tj ET\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = BytesIO(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(output.tell())
        output.write(f"{index} 0 obj\n".encode())
        output.write(obj)
        output.write(b"\nendobj\n")
    xref_offset = output.tell()
    output.write(f"xref\n0 {len(objects) + 1}\n".encode())
    output.write(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.write(f"{offset:010d} 00000 n \n".encode())
    output.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n"
        f"{xref_offset}\n%%EOF\n".encode()
    )
    return output.getvalue()


def test_extract_text_decodes_utf8_bom_txt_and_md():
    assert (
        extract_text("notes.txt", "\ufeff第一行\n第二行".encode()) == "第一行\n第二行"
    )
    assert extract_text("notes.md", "\ufeff# 标题\n\n正文".encode()) == "# 标题\n\n正文"


def test_extract_text_extracts_docx_paragraphs():
    document = DocxDocument()
    document.add_paragraph("第一段")
    document.add_paragraph("第二段")
    content = BytesIO()
    document.save(content)

    assert extract_text("brief.docx", content.getvalue()) == "第一段\n第二段"


def test_extract_text_extracts_pdf_text():
    assert extract_text("brief.pdf", make_pdf_with_text("PDF body")) == "PDF body"


@pytest.mark.parametrize(
    ("filename", "content"),
    [
        ("empty.pdf", b""),
        ("empty.txt", b"\xef\xbb\xbf\n\t"),
    ],
)
def test_extract_text_rejects_empty_or_unextractable_content(filename, content):
    with pytest.raises(AppError) as captured:
        extract_text(filename, content)

    assert captured.value.code == "DOCUMENT_CONTENT_NOT_EXTRACTABLE"
    assert captured.value.message == "文档未提取到可索引的正文"
