from collections.abc import Iterable
from html import escape
from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.models import CompetitorReport, TaskEvidence

_TYPES = ("web", "private_knowledge", "model_inference")
_LABELS = {
    "web": "网页证据",
    "private_knowledge": "私有知识库证据",
    "model_inference": "模型推断",
}


def build_pdf(report: CompetitorReport, evidence: Iterable[TaskEvidence]) -> bytes:
    buffer = BytesIO()
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "ChineseTitle", parent=styles["Title"], fontName="STSong-Light"
    )
    heading = ParagraphStyle(
        "ChineseHeading", parent=styles["Heading2"], fontName="STSong-Light"
    )
    body = ParagraphStyle(
        "ChineseBody", parent=styles["BodyText"], fontName="STSong-Light"
    )
    story: list[Any] = [Paragraph(_xml(report.title), title), Spacer(1, 8)]
    content = report.structured_json
    if summary := _text(content.get("executive_summary")):
        story.extend(
            [
                Paragraph("执行摘要", heading),
                Paragraph(_xml(summary), body),
                Spacer(1, 8),
            ]
        )
    competitors = _competitors(content)
    if competitors:
        story.extend(
            [Paragraph("竞品对比", heading), _table(competitors), Spacer(1, 8)]
        )
    recommendations = _recommendations(competitors)
    if recommendations:
        story.append(Paragraph("建议", heading))
        story.extend(Paragraph(f"• {_xml(item)}", body) for item in recommendations)
        story.append(Spacer(1, 8))
    _append_pdf_evidence(story, evidence, heading, body)
    document = SimpleDocTemplate(
        buffer, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm
    )
    document.build(
        story,
        onFirstPage=_footer(report.title),
        onLaterPages=_footer(report.title),
    )
    return buffer.getvalue()


def build_docx(report: CompetitorReport, evidence: Iterable[TaskEvidence]) -> bytes:
    document = Document()
    document.sections[0].left_margin = Inches(0.7)
    document.sections[0].right_margin = Inches(0.7)
    document.add_heading(report.title, level=0)
    content = report.structured_json
    if summary := _text(content.get("executive_summary")):
        document.add_heading("执行摘要", level=1)
        document.add_paragraph(summary)
    competitors = _competitors(content)
    if competitors:
        document.add_heading("竞品对比", level=1)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, value in zip(
            table.rows[0].cells, ("竞品", "定位", "核心功能", "优势"), strict=True
        ):
            cell.text = value
        for competitor in competitors:
            cells = table.add_row().cells
            cells[0].text = _text(competitor.get("name"))
            cells[1].text = _text(competitor.get("positioning"))
            cells[2].text = "、".join(_strings(competitor.get("core_features")))
            cells[3].text = "、".join(_strings(competitor.get("strengths")))
    if recommendations := _recommendations(competitors):
        document.add_heading("建议", level=1)
        for recommendation in recommendations:
            document.add_paragraph(recommendation, style="List Bullet")
    _append_docx_evidence(document, evidence)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _competitors(content: dict[str, object]) -> list[dict[str, object]]:
    value = content.get("competitors", [])
    return (
        [item for item in value if isinstance(item, dict)]
        if isinstance(value, list)
        else []
    )


def _recommendations(competitors: Iterable[dict[str, object]]) -> list[str]:
    return [
        item
        for competitor in competitors
        for item in _strings(competitor.get("recommendations"))
    ]


def _table(competitors: list[dict[str, object]]) -> Table:
    style = ParagraphStyle("TableCell", fontName="STSong-Light", fontSize=8, leading=11)
    rows = [
        [
            Paragraph(_xml(label), style)
            for label in ("竞品", "定位", "核心功能", "优势")
        ]
    ]
    rows.extend(
        [
            [
                Paragraph(_xml(_text(item.get("name"))), style),
                Paragraph(_xml(_text(item.get("positioning"))), style),
                Paragraph(_xml("、".join(_strings(item.get("core_features")))), style),
                Paragraph(_xml("、".join(_strings(item.get("strengths")))), style),
            ]
            for item in competitors
        ]
    )
    table = Table(rows, colWidths=[30 * mm, 45 * mm, 50 * mm, 39 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF8")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    return table


def _append_pdf_evidence(story, evidence, heading, body) -> None:
    story.append(Paragraph("证据链接", heading))
    cell = ParagraphStyle("EvidenceCell", parent=body, fontSize=8, leading=11)
    rows = [
        [Paragraph("证据类型", cell), Paragraph("来源", cell), Paragraph("链接", cell)]
    ]
    for source_type, entries in _group(evidence).items():
        if not entries:
            rows.append(
                [
                    Paragraph(_LABELS[source_type], cell),
                    Paragraph("暂无", cell),
                    Paragraph("", cell),
                ]
            )
        for entry in entries:
            url = _text(entry.source_url)
            link = _link(url) if url else "无链接"
            rows.append(
                [
                    Paragraph(_LABELS[source_type], cell),
                    Paragraph(_xml(entry.source_title), cell),
                    Paragraph(link, cell),
                ]
            )
    table = Table(rows, colWidths=[32 * mm, 58 * mm, 74 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF8")),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(table)


def _append_docx_evidence(document: Document, evidence: Iterable[TaskEvidence]) -> None:
    document.add_heading("证据链接", level=1)
    for source_type, entries in _group(evidence).items():
        document.add_heading(_LABELS[source_type], level=2)
        if not entries:
            document.add_paragraph("暂无")
        for entry in entries:
            document.add_paragraph(
                f"{entry.source_title}: {entry.source_url or '无链接'}",
                style="List Bullet",
            )


def _group(evidence: Iterable[TaskEvidence]) -> dict[str, list[TaskEvidence]]:
    grouped = {source_type: [] for source_type in _TYPES}
    for entry in evidence:
        if entry.source_type in grouped:
            grouped[entry.source_type].append(entry)
    return grouped


def _strings(value: object) -> list[str]:
    return (
        [item for item in value if isinstance(item, str)]
        if isinstance(value, list)
        else []
    )


def _text(value: object) -> str:
    return value if isinstance(value, str) else ""


def _xml(value: str) -> str:
    return escape(value, quote=True)


def _link(url: str) -> str:
    escaped = _xml(url)
    return f'<link href="{escaped}">{escaped}</link>'


def _footer(title: str):
    safe_title = title.replace("\r", " ").replace("\n", " ")[:80]

    def draw_footer(canvas, document) -> None:
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.drawString(18 * mm, 12 * mm, safe_title)
        canvas.drawRightString(A4[0] - 18 * mm, 12 * mm, f"第 {document.page} 页")
        canvas.restoreState()

    return draw_footer
